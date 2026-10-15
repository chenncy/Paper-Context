"""Offline MCP end-to-end check, using temporary paper and memory databases."""
from __future__ import annotations

import asyncio
import argparse
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
USERNAME = "smoke-user"
MEMORY_KEY = "reading_preference"
MEMORY_VALUE = "Always include the paper title, section and page when citing evidence."


def isolated_env(directory: Path) -> dict[str, str]:
    env = dict(os.environ)
    for key in list(env):
        if key.upper().startswith(("MYSQL_", "EMBEDDING_", "DASHSCOPE_", "OPENAI_")):
            env.pop(key)
    env.update({
        "PAPER_CONTEXT_SKIP_DOTENV": "1",
        "PYTHONIOENCODING": "utf-8",
        "SQLALCHEMY_DATABASE_URL": f"sqlite:///{(directory / 'memory.db').as_posix()}",
        "DB_DIR": str(directory / "vectors"),
        "STORE_DIR": str(directory / "parents"),
        "PAPERS_DIR": str(directory / "papers"),
        "RAG_FTS_PATH": str(directory / "fts.db"),
        "VECTOR_BACKEND": "chroma",
        "VECTOR_COLLECTION_NAME": "smoke_papers",
        "EPISODIC_MEMORY_COLLECTION": "smoke_episodic",
        "GLOBAL_PROFILE_SESSION_ID": "GLOBAL_PROFILE",
        "RAG_RETRIEVAL_MODE": "auto",
        "PDF_PARSER_MODE": "layout",
        "PDF_PARSER_OCR_ENABLED": "false",
        "RAG_RERANK_ENABLED": "false",
        "RAG_CONTEXTUAL_ENABLED": "false",
        "RAG_VLM_ENABLED": "false",
        "COLLECTION_REVISION_ENABLED": "false",
        "EMBEDDING_CACHE_ENABLED": "false",
        "REDIS_URL": "",
        "RETRIEVAL_CACHE_ENABLED": "false",
        "EPISODIC_MEMORY_ENABLED": "false",
        "MEMORY_GRAPH_ENABLED": "false",
        "ANONYMIZED_TELEMETRY": "FALSE",
        "RAG_MCP_TRANSPORT": "stdio",
        "MEMORY_MCP_TRANSPORT": "stdio",
    })
    return env


async def call(session, name: str, arguments: dict):
    result = await session.call_tool(name, arguments)
    if result.isError:
        raise RuntimeError(f"{name}: MCP reported an error")
    payload = "\n".join(item.text for item in result.content if item.type == "text")
    try:
        value = json.loads(payload)
    except json.JSONDecodeError:
        return payload
    if isinstance(value, dict) and (value.get("error") or value.get("success") is False):
        raise RuntimeError(f"{name}: {value.get('error', 'operation failed')}")
    return value


async def check_service(service: str, directory: Path, env: dict, *, fresh: bool = False) -> None:
    from mcp import ClientSession, StdioServerParameters
    from mcp.client.stdio import stdio_client

    parameters = StdioServerParameters(
        command=sys.executable,
        args=[str(ROOT / f"{service}_mcp_server.py")],
        # Start outside the repository to verify absolute entry points.
        cwd=str(directory),
        env=env,
    )
    with (directory / f"{service}.log").open("w", encoding="utf-8") as log:
        async with stdio_client(parameters, errlog=log) as (read, write):
            async with ClientSession(read, write) as session:
                await session.initialize()
                tools = {tool.name for tool in (await session.list_tools()).tools}
                expected = {"rag_ingest", "rag_search", "rag_stats"} if service == "rag" else {
                    "memory_save", "memory_get", "memory_prepare",
                }
                if not expected <= tools:
                    raise RuntimeError(f"{service}: required tools missing")
                print(f"[PASS] {service} MCP startup and tool discovery", flush=True)
                if service == "rag":
                    for suffix, query in (("txt", "retrieval provenance"), ("docx", "document evidence"), ("pdf", "academic retrieval")):
                        filename = directory / f"demo.{suffix}"
                        ingested = await call(session, "rag_ingest", {
                            "file_path": str(filename), "paper_title": f"Demo {suffix.upper()}",
                        })
                        if not ingested["success"] or ingested["validation_metrics"]["vectorized"]:
                            raise RuntimeError(f"{suffix}: expected successful offline ingestion")
                        for mode in ("fts", "auto"):
                            found = await call(session, "rag_search", {"query": query, "mode": mode})
                            hits = [h for h in found["hits"] if h["source"] == filename.name]
                            if not hits or found["mode"] != "fts" or not hits[0]["title"] or not hits[0]["section"] or hits[0]["page"] is None:
                                raise RuntimeError(f"{suffix}: missing evidence/provenance in {mode} search")
                        print(f"[PASS] {suffix.upper()} ingestion -> FTS/auto search -> title/section/page", flush=True)
                    stats = await call(session, "rag_stats", {})
                    if stats["fts_chunk_count"] <= 0 or stats.get("vector_status") == "unavailable":
                        raise RuntimeError("rag_stats: local backend unavailable")
                    print("[PASS] local index statistics", flush=True)
                elif not fresh:
                    saved = await call(session, "memory_save", {
                        "username": USERNAME, "key": MEMORY_KEY, "value": MEMORY_VALUE,
                        "scope": "global", "session_id": "first-session",
                    })
                    if saved["status"] != "ok":
                        raise RuntimeError("memory_save failed")
                    print("[PASS] global memory saved in first session", flush=True)
                else:
                    read_value = await call(session, "memory_get", {
                        "username": USERNAME, "key": MEMORY_KEY, "session_id": "new-session",
                    })
                    prepared = await call(session, "memory_prepare", {
                        "username": USERNAME, "session_id": "new-session",
                    })
                    if MEMORY_VALUE not in read_value or MEMORY_VALUE not in prepared["content"]:
                        raise RuntimeError("memory did not survive a new session/server process")
                    print("[PASS] new server/session -> memory_get and memory_prepare recall saved preference", flush=True)


async def scenario(directory: Path, env: dict) -> None:
    await check_service("rag", directory, env)
    await check_service("memory", directory, env)
    await check_service("memory", directory, env, fresh=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--timeout", type=int, default=300, help="MCP scenario timeout in seconds")
    args = parser.parse_args()
    if args.timeout <= 0:
        parser.error("--timeout must be positive")
    # Suppress Windows crash dialogs in this process and inherited native children.
    if os.name == "nt":
        import ctypes
        ctypes.windll.kernel32.SetErrorMode(3)
    try:
        with tempfile.TemporaryDirectory(prefix="paper-context-smoke-") as temporary:
            directory = Path(temporary)
            env = isolated_env(directory)
            initialized = subprocess.run(
                [sys.executable, str(ROOT / "scripts/init_db.py")],
                cwd=directory, env=env, capture_output=True, timeout=60,
            )
            if initialized.returncode:
                raise RuntimeError("temporary SQLite initialization failed")
            print("[PASS] temporary SQLite initialization", flush=True)
            (directory / "demo.txt").write_text(
                (ROOT / "examples/demo-paper.txt").read_text(encoding="utf-8"), encoding="utf-8",
            )
            from docx import Document
            document = Document()
            document.add_paragraph("# Introduction")
            document.add_paragraph("Document evidence is retrieved from a Word file. This local study preserves source metadata and section boundaries for later verification.")
            document.save(directory / "demo.docx")
            import pymupdf
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "1 Introduction", fontsize=16)
                page.insert_text((72, 105), "Academic retrieval preserves evidence for local research.", fontsize=11)
                pdf.save(str(directory / "demo.pdf"))
            asyncio.run(asyncio.wait_for(scenario(directory, env), timeout=args.timeout))
        print("[PASS] offline end-to-end check complete; existing user data was not used", flush=True)
        return 0
    except Exception as exc:
        print(f"[FAIL] {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
