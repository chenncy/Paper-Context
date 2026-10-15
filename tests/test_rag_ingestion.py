"""入库回归：真实解析/切分/本地索引，embedding 使用本地确定性替身。"""
import json
import sqlite3
import sys

import pytest
from langchain_core.embeddings import DeterministicFakeEmbedding, Embeddings

from app import collection_revision, config, embedding_cache, rag_fts, rag_service
from app import retriever, retrieval_cache, vector_store
from app.ingestion import pipeline
from app.ingestion.pipeline import run_ingestion_pipeline


@pytest.fixture
def isolated_ingestion(tmp_path, monkeypatch):
    # 已从开源仓库移除的 worker 必须完全不参与入库。
    monkeypatch.setitem(sys.modules, "app.embedding_worker", None)
    monkeypatch.setattr(collection_revision, "COLLECTION_REVISION_ENABLED", False)
    monkeypatch.setattr(config, "RAG_CONTEXTUAL_ENABLED", False)
    monkeypatch.setattr(config, "RAG_VLM_ENABLED", False)
    monkeypatch.setattr(vector_store, "VECTOR_BACKEND", "chroma")
    monkeypatch.setattr(retriever, "VECTOR_BACKEND", "chroma")
    monkeypatch.setattr(vector_store, "DB_DIR", str(tmp_path / "vectors"))
    monkeypatch.setattr(retriever, "STORE_DIR", str(tmp_path / "parents"))
    monkeypatch.setattr(config, "STORE_DIR", str(tmp_path / "parents"))
    monkeypatch.setattr(retrieval_cache, "get_redis_client", lambda: None)
    monkeypatch.setenv("ANONYMIZED_TELEMETRY", "FALSE")
    rag_fts.reset_for_tests(str(tmp_path / "fts.db"))
    rag_service.invalidate_retriever_cache()
    path = tmp_path / "paper.txt"
    path.write_text(
        "# Introduction\n"
        "A retrieval augmented generation system keeps evidence traceable to its "
        "source. This study examines a local knowledge base for academic papers.\n\n"
        "# Method\n"
        "The method stores paragraph evidence and records the source and section "
        "for every chunk. Full text search retrieves relevant evidence locally.\n",
        encoding="utf-8",
    )
    yield path
    rag_service.invalidate_retriever_cache()
    rag_fts.reset_for_tests(str(tmp_path / "cleanup.db"))


@pytest.fixture
def local_ingestion(isolated_ingestion, monkeypatch):
    monkeypatch.setattr(
        embedding_cache, "CachedOpenAIEmbeddings",
        lambda **kwargs: DeterministicFakeEmbedding(size=8),
    )
    return isolated_ingestion


@pytest.mark.parametrize("paper_id", [None, 73])
def test_ingest_search_delete_without_embedding_worker(local_ingestion, paper_id):
    result = rag_service.rag_ingest(
        str(local_ingestion), paper_id=paper_id, paper_title="Local evidence study",
    )
    assert result["success"], result
    assert result["validation_metrics"]["vectorized"] is True
    assert result["validation_metrics"]["chunk_count"] > 0
    hits = rag_service.rag_search("retrieval augmented generation", mode="fts")["hits"]
    assert hits and hits[0]["source"] == "paper.txt"
    assert hits[0]["title"] == "Local evidence study"
    assert hits[0]["section"]
    counts = rag_fts.counts()
    assert counts["fts_chunk_count"] > 0
    assert counts["not_vectorized"] == 0

    deleted = rag_service.rag_ingest("", **result["reverse"])
    assert deleted["success"]
    assert deleted["deleted_chunks"] == counts["fts_chunk_count"]
    assert deleted["deleted_fts_chunks"] == counts["fts_chunk_count"]
    assert rag_fts.counts()["fts_chunk_count"] == 0
    assert rag_service.rag_search("retrieval", mode="fts")["hits"] == []


def test_parse_without_index_does_not_import_worker(local_ingestion):
    result = run_ingestion_pipeline(str(local_ingestion), skip_index=True)
    assert result.success, result.logs
    assert result.validation["metrics"]["chunk_count"] > 0
    assert rag_fts.counts()["fts_chunk_count"] == 0


def test_unsupported_file_returns_ingestion_error(local_ingestion):
    path = local_ingestion.with_suffix(".unsupported")
    path.write_text("unsupported content", encoding="utf-8")
    result = rag_service.rag_ingest(str(path))
    assert result["success"] is False
    assert "不支持的文件格式" in result["error"]


class OfflineEmbeddings(Embeddings):
    def embed_documents(self, texts):
        raise ConnectionError("Connection error.")

    def embed_query(self, text):
        raise ConnectionError("Connection error.")


@pytest.mark.parametrize("paper_id", [None, 73])
@pytest.mark.parametrize("availability", ["no_key", "offline"])
def test_fts_only_ingest_search_stats_delete(isolated_ingestion, monkeypatch, availability, paper_id):
    if availability == "no_key":
        monkeypatch.setattr(retriever, "EMBEDDING_API_KEY", "")
        for key in ("OPENAI_API_KEY", "OPENAI_ADMIN_KEY", "DASHSCOPE_API_KEY", "EMBEDDING_API_KEY"):
            monkeypatch.delenv(key, raising=False)
    else:
        monkeypatch.setattr(embedding_cache, "CachedOpenAIEmbeddings", lambda **kwargs: OfflineEmbeddings())
    monkeypatch.setattr(pipeline.time, "sleep", lambda seconds: None)

    result = rag_service.rag_ingest(str(isolated_ingestion), paper_id=paper_id, paper_title="FTS study")
    assert result["success"], result
    assert result["validation_metrics"]["vectorized"] is False
    pending = list(rag_fts.iter_unvectorized())
    assert pending, "success must mean text was actually persisted"
    ids = [item["chunk_id"] for item in pending]
    assert all(item["metadata"].get("doc_id") for item in pending)
    for mode in ("fts", "auto"):
        search = rag_service.rag_search("retrieval augmented generation", mode=mode)
        assert search["mode"] == "fts"
        assert search["hits"] and search["hits"][0]["title"] == "FTS study"

    repeated = rag_service.rag_ingest(str(isolated_ingestion), paper_id=paper_id, paper_title="FTS study")
    assert repeated["success"]
    assert [item["chunk_id"] for item in rag_fts.iter_unvectorized()] == ids

    def no_embedding_initialization(**kwargs):
        raise AssertionError("stats and delete must not initialize an embedding client")
    monkeypatch.setattr(embedding_cache, "CachedOpenAIEmbeddings", no_embedding_initialization)
    import rag_mcp_server
    stats = json.loads(rag_mcp_server.rag_stats())
    assert "error" not in stats, stats
    assert stats["fts_chunk_count"] == len(pending)
    assert stats["chunk_count"] == 0
    assert stats["parent_doc_count"] > 0
    deleted = rag_service.rag_ingest("", **result["reverse"])
    assert deleted["success"], deleted
    assert deleted["deleted_fts_chunks"] == len(pending)
    assert deleted["deleted_chunks"] == 0
    assert rag_fts.counts()["fts_chunk_count"] == 0
    assert rag_service.rag_search("retrieval", mode="fts")["hits"] == []


def test_backfill_restores_parent_retrieval(isolated_ingestion, monkeypatch):
    monkeypatch.setattr(embedding_cache, "CachedOpenAIEmbeddings", lambda **kwargs: OfflineEmbeddings())
    monkeypatch.setattr(pipeline.time, "sleep", lambda seconds: None)
    result = rag_service.rag_ingest(str(isolated_ingestion), paper_id=73, paper_title="Backfill study")
    assert result["success"], result
    pending = list(rag_fts.iter_unvectorized())
    assert pending
    assert all(item["metadata"]["paper_id"] == 73 for item in pending)

    monkeypatch.setattr(
        embedding_cache, "CachedOpenAIEmbeddings",
        lambda **kwargs: DeterministicFakeEmbedding(size=8),
    )
    from scripts.backfill_embeddings import _flush
    vectorstore, _, _, _ = retriever._get_components()
    assert _flush(vectorstore, pending) == len(pending)
    assert rag_fts.counts()["not_vectorized"] == 0
    search = rag_service.rag_search("retrieval augmented generation", mode="vector")
    assert search["hits"], search
    assert all(hit["title"] == "Backfill study" for hit in search["hits"])
    assert all(not hit["text"].startswith("[Backfill study") for hit in search["hits"])
    deleted = rag_service.rag_ingest("", **result["reverse"])
    assert deleted["success"]
    assert deleted["deleted_chunks"] == len(pending)
    assert rag_fts.counts()["fts_chunk_count"] == 0


def test_fts_write_failure_is_not_reported_as_success(local_ingestion, monkeypatch):
    def fail_fts(*args, **kwargs):
        raise sqlite3.OperationalError("disk full")
    monkeypatch.setattr(rag_fts, "upsert_chunks", fail_fts)
    result = rag_service.rag_ingest(str(local_ingestion))
    assert result["success"] is False, result
    assert "disk full" in result["error"]


def test_delete_vector_failure_reports_partial_result(local_ingestion, monkeypatch):
    result = rag_service.rag_ingest(str(local_ingestion))
    assert result["success"]
    expected = rag_fts.counts()["fts_chunk_count"]
    def fail_vector_delete(source):
        raise ConnectionError("vector backend unavailable")
    monkeypatch.setattr(retriever, "delete_documents_by_source", fail_vector_delete)
    deleted = rag_service.rag_ingest("", **result["reverse"])
    assert deleted["success"] is False
    assert deleted["deleted_fts_chunks"] == expected
    assert "vector backend unavailable" in deleted["error"]
    assert rag_fts.counts()["fts_chunk_count"] == 0


def test_stats_keep_fts_counts_when_vector_backend_is_unavailable(local_ingestion, monkeypatch):
    assert rag_service.rag_ingest(str(local_ingestion))["success"]
    expected = rag_fts.counts()["fts_chunk_count"]
    def unavailable(*args, **kwargs):
        raise ConnectionError("vector backend unavailable")
    monkeypatch.setattr(retriever, "_get_vectorstore", unavailable)
    stats = rag_service.rag_stats()
    assert stats["fts_chunk_count"] == expected
    assert stats["chunk_count"] is None
    assert stats["vector_status"] == "unavailable"
    assert "vector backend unavailable" in stats["vector_error"]


def test_vector_write_without_ids_keeps_text_on_failure(isolated_ingestion):
    store = vector_store.create_vectorstore("no_ids", OfflineEmbeddings())
    with pytest.raises(ConnectionError):
        store.add_texts(["Evidence survives offline embedding."], metadatas=[{"source": "no_ids.txt"}])
    pending = list(rag_fts.iter_unvectorized())
    assert len(pending) == 1
    assert pending[0]["text"] == "Evidence survives offline embedding."
    assert pending[0]["metadata"]["source"] == "no_ids.txt"


def test_real_pdf_fts_only_without_key(isolated_ingestion, monkeypatch):
    import fitz
    monkeypatch.setattr(retriever, "EMBEDDING_API_KEY", "")
    for key in ("OPENAI_API_KEY", "OPENAI_ADMIN_KEY"):
        monkeypatch.delenv(key, raising=False)
    path = isolated_ingestion.with_suffix(".pdf")
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((72, 72), "1 Introduction", fontsize=16)
        page.insert_text((72, 105), "Retrieval evidence survives offline ingestion for local academic research.", fontsize=11)
        pdf.save(str(path))
    result = rag_service.rag_ingest(str(path), paper_title="PDF evidence study")
    assert result["success"], result
    assert result["validation_metrics"]["vectorized"] is False
    search = rag_service.rag_search("retrieval evidence", mode="fts")
    assert search["hits"]
    assert search["hits"][0]["source"] == "paper.pdf"
    assert search["hits"][0]["page"] == "0"
    assert rag_service.rag_ingest("", **result["reverse"])["success"]
    assert rag_fts.counts()["fts_chunk_count"] == 0


def test_real_docx_fts_only_without_key(isolated_ingestion, monkeypatch):
    from docx import Document
    monkeypatch.setattr(retriever, "EMBEDDING_API_KEY", "")
    for key in ("OPENAI_API_KEY", "OPENAI_ADMIN_KEY"):
        monkeypatch.delenv(key, raising=False)
    path = isolated_ingestion.with_suffix(".docx")
    document = Document()
    document.add_paragraph("# Introduction")
    document.add_paragraph("Document evidence survives offline ingestion for academic research. The local knowledge base preserves the source and section of each paragraph.")
    document.save(path)
    result = rag_service.rag_ingest(str(path), paper_title="DOCX evidence study")
    assert result["success"], result
    assert result["validation_metrics"]["vectorized"] is False
    search = rag_service.rag_search("document evidence", mode="fts")
    assert search["hits"] and search["hits"][0]["source"] == "paper.docx"
    assert search["hits"][0]["title"] == "DOCX evidence study"
    assert search["hits"][0]["section"]
    assert rag_service.rag_ingest("", **result["reverse"])["success"]
    assert rag_fts.counts()["fts_chunk_count"] == 0


def test_mcp_wrappers_do_not_print_diagnostics_to_stdout(local_ingestion, capsys):
    import rag_mcp_server
    result = json.loads(rag_mcp_server.rag_ingest(str(local_ingestion)))
    assert result["success"]
    assert json.loads(rag_mcp_server.rag_search("retrieval", mode="auto"))["hits"]
    assert json.loads(rag_mcp_server.rag_stats())["fts_chunk_count"] > 0
    assert json.loads(rag_mcp_server.rag_ingest("", **result["reverse"]))["success"]
    assert capsys.readouterr().out == "", "stdio stdout is reserved for JSON-RPC messages"
